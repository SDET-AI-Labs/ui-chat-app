"""API test case insertion with Config3 constraints and duplicate prevention."""
from typing import Dict, Any, List
from pathlib import Path
import openpyxl
import json
import hashlib
from datetime import datetime

from . import run_excel_suite

# Configuration constants
MAX_INSERTS_PER_REQUEST = 3
API_TEST_FILE = "./data/api_tests.xlsx"
API_TEST_SHEET = "api_tests"

class DuplicateTestError(Exception):
    """Raised when attempting to insert a duplicate test case."""
    pass

class InsertLimitError(Exception):
    """Raised when insert request exceeds Config3 limit."""
    pass

def _normalize_empty_body(body: Any) -> str:
    """Normalize different forms of empty body to canonical empty string."""
    if body is None:
        return ""
    body_str = str(body).strip()
    if body_str in ["", "{}", "null"]:
        return ""
    return body_str

def _generate_test_key(test_case: Dict[str, Any]) -> str:
    """Generate uniqueness key for test case.
    
    Key = normalized(method|url|body)
    where:
    - method is uppercased and stripped
    - url has trailing slashes removed
    - body is normalized to canonical empty string
    """
    # Normalize method: uppercase and strip
    method = str(test_case.get('method', '')).strip().upper()
    
    # Normalize URL: strip and remove trailing slash
    url = str(test_case.get('url', '')).strip().rstrip('/')
    
    # Normalize body: convert None/empty/"{}" to canonical empty string
    body = _normalize_empty_body(test_case.get('body'))
    
    # Always include all parts for consistent key generation
    key_str = f"{method}|{url}|{body}"
    return hashlib.sha256(key_str.encode()).hexdigest()

def _load_existing_keys(wb: openpyxl.Workbook, sheet: str) -> set:
    """Load uniqueness keys for existing test cases."""
    ws = wb[sheet]
    headers = [str(cell.value).lower() for cell in ws[1]]
    
    # Find required column indices
    method_idx = headers.index('method')
    url_idx = headers.index('url')
    body_idx = headers.index('body')
    
    existing_keys = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        test_case = {
            'method': row[method_idx],
            'url': row[url_idx],
            'body': row[body_idx]
        }
        existing_keys.add(_generate_test_key(test_case))
    
    return existing_keys

async def run(args: Dict[str, Any]) -> Dict[str, Any]:
    """Insert API test cases with duplicate prevention and run tests.
    
    Args:
        test_cases: List of test cases to insert, each containing:
            - module: Test module/group
            - description: Test description 
            - method: HTTP method
            - url: API endpoint
            - body: Request body (optional)
            - headers: Request headers (optional)
            - expected_status: Expected HTTP status
            - expected_field: Field to check in response
            - expected_value: Expected value
            
    Returns:
        Dict with test execution results after insertion:
        {
            "passed": int,
            "failed": int,
            "skipped": int,
            "results": List[Dict]
        }
        
    Raises:
        InsertLimitError: If test_cases exceeds MAX_INSERTS_PER_REQUEST
        DuplicateTestError: If any test case already exists
        ValueError: For invalid/missing required fields
    """
    test_cases = args.get('test_cases', [])
    
    # Validate insert limit
    if len(test_cases) > MAX_INSERTS_PER_REQUEST:
        raise InsertLimitError(
            f"Maximum {MAX_INSERTS_PER_REQUEST} test cases allowed per request. "
            f"Got {len(test_cases)}"
        )
    
    # Required fields for each test case
    required_fields = {
        'module', 'description', 'method', 'url', 
        'expected_status', 'expected_field', 'expected_value'
    }
    
    # Load workbook
    file_path = Path(API_TEST_FILE)
    if not file_path.exists():
        raise ValueError(f"API test file not found: {API_TEST_FILE}")
        
    wb = openpyxl.load_workbook(file_path)
    if API_TEST_SHEET not in wb.sheetnames:
        raise ValueError(f"Sheet {API_TEST_SHEET} not found")
        
    ws = wb[API_TEST_SHEET]
    
    # Get existing test case keys
    existing_keys = _load_existing_keys(wb, API_TEST_SHEET)
    
    # Validate all test cases before insertion
    for test in test_cases:
        # Check required fields
        missing = required_fields - set(test.keys())
        if missing:
            raise ValueError(f"Missing required fields: {missing}")
            
        # Check for duplicates
        test_key = _generate_test_key(test)
        if test_key in existing_keys:
            raise DuplicateTestError(
                f"Duplicate test case: {test['method']} {test['url']}"
            )
    
    # All validation passed, insert test cases
    headers = [str(cell.value).lower() for cell in ws[1]]
    
    for test in test_cases:
        # Prepare row data
        row_data = []
        for header in headers:
            if header == 'status':
                row_data.append('active')  # Default status
            elif header == 'last_run':
                row_data.append(datetime.now().strftime('%Y-%m-%d %H:%M'))
            elif header == 'result':
                row_data.append('')  # Empty initial result
            elif header == 'actual_status':
                row_data.append('')  # Empty actual values
            elif header == 'actual_value':
                row_data.append('')
            else:
                row_data.append(test.get(header, ''))
        
        ws.append(row_data)
    
    # Save changes
    wb.save(file_path)
    wb.close()
    
    # Run tests and return results
    return await run_excel_suite.run({
        'file': API_TEST_FILE,
        'sheet': API_TEST_SHEET
    })