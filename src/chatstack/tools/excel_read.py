from __future__ import annotations

import os
import asyncio
from typing import Any, Dict, List
import openpyxl


def _read_excel_sync(abs_path: str, sheet: str) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(abs_path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {sheet}")
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h) if h is not None else f"col{i+1}" for i, h in enumerate(rows[0])]
        result: List[Dict[str, Any]] = []
        for row in rows[1:]:
            row_dict = {headers[i]: row[i] for i in range(len(headers))}
            result.append(row_dict)
        return result
    finally:
        try:
            wb.close()
        except Exception:
            pass


async def run(args: Dict[str, Any]) -> Dict[str, Any]:
    """
    Args:
        args: {
            "path": "<xlsx file under ./data/>",
            "sheet": "<sheet name>"
        }
    Returns:
        dict: {"result": <list of dicts>} or {"result": "error", "error": <str>}
    """
    try:
        path = args.get("path") if isinstance(args, dict) else None
        sheet = args.get("sheet") if isinstance(args, dict) else None
        if not isinstance(path, str) or not path.strip():
            return {"result": "error", "error": "Missing or invalid 'path' argument"}
        if not isinstance(sheet, str) or not sheet.strip():
            return {"result": "error", "error": "Missing or invalid 'sheet' argument"}

        # Resolve and enforce ./data sandbox
        base = os.path.abspath("./data")
        abs_path = os.path.abspath(path)
        if not abs_path.startswith(base + os.sep) and abs_path != base:
            return {"result": "error", "error": "File must be under ./data/ directory"}
        if not os.path.exists(abs_path):
            return {"result": "error", "error": f"File not found: {path}"}

        rows = await asyncio.to_thread(_read_excel_sync, abs_path, sheet)
        return {"result": rows}
    except Exception as e:
        return {"result": "error", "error": str(e)}
